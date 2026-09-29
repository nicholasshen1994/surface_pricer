"""Term sheets for real trades: loading and validation.

A term sheet is a flat record (JSON object, CSV row or Excel row) describing
one contract.  Vanilla trades are fully supported today; exotic product types
are routed through :mod:`surface_pricer.pricing.exotics` once a pricer is
registered there.

Example (JSON)::

    {"trades": [
      {"trade_id": "TRD-001", "underlying": "MO", "product_type": "vanilla",
       "booked_date": "2026-06-01", "start_date": "2026-06-01",
       "expiry_date": "2027-06-18", "call_put": "call", "strike": 7500.0,
       "strike_type": "absolute", "notional": 1000000,
       "ki_barrier": null, "ko_barrier": null, "ki_flag": false, "ko_flag": false,
       "coupon": null, "observations": [{"date": "2026-06-18", "spot": 7480.0}]}
    ]}

The same fields work as CSV / Excel columns; ``observations`` is then a JSON
array encoded in a single cell.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

from ..core.daycount import DateLike, to_datetime

SUPPORTED_PRODUCT_TYPES: Tuple[str, ...] = ("vanilla",)
_VANILLA_ALIASES = {"vanilla", "van", "european", "euro_vanilla"}

_TRUE_TEXT = {"1", "true", "t", "yes", "y"}
_FALSE_TEXT = {"0", "false", "f", "no", "n"}


@dataclass(frozen=True)
class ObservationRecord:
    """One recorded observation of a barrier / autocall monitoring date."""

    date: DateLike
    spot: float

    def __post_init__(self):
        object.__setattr__(self, "date", to_datetime(self.date))
        object.__setattr__(self, "spot", float(self.spot))


@dataclass
class TradeTerms:
    """Flat description of one trade (fields mirror the term-sheet columns)."""

    trade_id: str
    underlying: str
    product_type: str = "vanilla"
    booked_date: Optional[DateLike] = None
    start_date: Optional[DateLike] = None
    expiry_date: Optional[DateLike] = None
    call_put: str = "call"
    strike: float = 0.0
    strike_type: str = "absolute"
    notional: float = 1.0
    ki_barrier: Optional[float] = None
    ko_barrier: Optional[float] = None
    ki_flag: bool = False
    ko_flag: bool = False
    coupon: Optional[float] = None
    observations: Tuple[ObservationRecord, ...] = ()
    metadata: Dict[str, Any] = field(default_factory=dict)

    def __post_init__(self):
        self.trade_id = str(self.trade_id or "").strip()
        if not self.trade_id:
            raise ValueError("trade_id must not be empty")
        self.underlying = str(self.underlying or "").strip().upper()
        if not self.underlying:
            raise ValueError("underlying must not be empty for {}".format(self.trade_id))
        self.product_type = str(self.product_type or "vanilla").strip().lower()
        for name in ("booked_date", "start_date", "expiry_date"):
            value = getattr(self, name)
            if value is not None:
                setattr(self, name, to_datetime(value))
        if self.expiry_date is None:
            raise ValueError("expiry_date is required for {}".format(self.trade_id))
        if self.start_date is not None and self.start_date > self.expiry_date:
            raise ValueError(
                "start_date {} is after expiry_date {} for {}".format(
                    self.start_date.date(), self.expiry_date.date(), self.trade_id
                )
            )
        self.call_put = str(self.call_put or "call").strip().lower()
        if self.call_put in {"c"}:
            self.call_put = "call"
        if self.call_put in {"p"}:
            self.call_put = "put"
        if self.call_put not in {"call", "put"}:
            raise ValueError(
                "call_put must be call or put for {}, got {!r}".format(
                    self.trade_id, self.call_put
                )
            )
        self.strike_type = str(self.strike_type or "absolute").strip().lower()
        self.strike = float(self.strike)
        if self.strike <= 0.0:
            raise ValueError("strike must be positive for {}".format(self.trade_id))
        self.notional = float(self.notional)
        for name in ("ki_barrier", "ko_barrier", "coupon"):
            value = getattr(self, name)
            if value is not None:
                setattr(self, name, float(value))
        self.ki_flag = bool(self.ki_flag)
        self.ko_flag = bool(self.ko_flag)
        self.observations = tuple(self.observations)

    @property
    def is_vanilla(self) -> bool:
        return self.product_type in _VANILLA_ALIASES

    def to_dict(self) -> Dict[str, Any]:
        payload: Dict[str, Any] = {
            "trade_id": self.trade_id,
            "underlying": self.underlying,
            "product_type": self.product_type,
            "call_put": self.call_put,
            "strike": self.strike,
            "strike_type": self.strike_type,
            "notional": self.notional,
            "ki_flag": self.ki_flag,
            "ko_flag": self.ko_flag,
        }
        for name in ("booked_date", "start_date", "expiry_date"):
            value = getattr(self, name)
            if value is not None:
                payload[name] = value.isoformat(sep=" ")
        for name in ("ki_barrier", "ko_barrier", "coupon"):
            value = getattr(self, name)
            if value is not None:
                payload[name] = value
        if self.observations:
            payload["observations"] = [
                {"date": item.date.isoformat(sep=" "), "spot": item.spot}
                for item in self.observations
            ]
        if self.metadata:
            payload["metadata"] = dict(self.metadata)
        return payload

    @staticmethod
    def from_dict(payload: Mapping[str, Any]) -> "TradeTerms":
        if not isinstance(payload, Mapping):
            raise ValueError("each trade must be a mapping, got {!r}".format(type(payload).__name__))
        return TradeTerms(
            trade_id=payload.get("trade_id"),
            underlying=payload.get("underlying"),
            product_type=payload.get("product_type", "vanilla"),
            booked_date=_clean(payload.get("booked_date")),
            start_date=_clean(payload.get("start_date")),
            expiry_date=_clean(payload.get("expiry_date", payload.get("expiry"))),
            call_put=payload.get("call_put", payload.get("option_type", "call")),
            strike=_clean(payload.get("strike")) or 0.0,
            strike_type=payload.get("strike_type", "absolute"),
            notional=_clean(payload.get("notional")) or 1.0,
            ki_barrier=_clean(payload.get("ki_barrier")),
            ko_barrier=_clean(payload.get("ko_barrier")),
            ki_flag=_to_bool(payload.get("ki_flag")),
            ko_flag=_to_bool(payload.get("ko_flag")),
            coupon=_clean(payload.get("coupon")),
            observations=_parse_observations(payload.get("observations")),
            metadata=dict(payload.get("metadata") or {}),
        )


def _clean(value: Any) -> Any:
    """Treat blank CSV cells as missing."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text or text.upper() in {"NAN", "NONE", "NULL", "-"}:
            return None
        return text
    return value


def _to_bool(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in _TRUE_TEXT:
        return True
    if text in _FALSE_TEXT or not text:
        return False
    raise ValueError("cannot interpret {!r} as a boolean flag".format(value))


def _parse_observations(value: Any) -> Tuple[ObservationRecord, ...]:
    if not value:
        return ()
    payload = value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return ()
        payload = json.loads(text)
    if not isinstance(payload, Sequence):
        raise ValueError("observations must be a list of {date, spot} objects")
    records: List[ObservationRecord] = []
    for item in payload:
        if not isinstance(item, Mapping) or "date" not in item or "spot" not in item:
            raise ValueError("each observation needs a date and a spot")
        records.append(ObservationRecord(date=item["date"], spot=float(item["spot"])))
    return tuple(sorted(records, key=lambda record: record.date))


# ------------------------------------------------------------------- loaders
def load_terms(path: str) -> List[TradeTerms]:
    """Load a term sheet from JSON / CSV / Excel, dispatching on the suffix."""
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError("term sheet not found: {}".format(path))
    suffix = source.suffix.lower()
    if suffix == ".json":
        payloads = _load_json(source)
    elif suffix in {".csv", ".txt"}:
        payloads = _load_csv(source)
    elif suffix in {".xlsx", ".xlsm"}:
        payloads = _load_xlsx(source)
    else:
        raise ValueError(
            "unsupported term-sheet format {!r}; use .json, .csv or .xlsx".format(suffix)
        )
    trades = [TradeTerms.from_dict(item) for item in payloads]
    if not trades:
        raise ValueError("term sheet {} contains no trades".format(path))
    seen = set()
    for trade in trades:
        if trade.trade_id in seen:
            raise ValueError("duplicate trade_id {!r} in {}".format(trade.trade_id, path))
        seen.add(trade.trade_id)
    return trades


def _load_json(source: Path) -> List[Mapping[str, Any]]:
    with open(source, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if isinstance(payload, Mapping):
        payload = payload.get("trades", [payload])
    if not isinstance(payload, Sequence):
        raise ValueError("{}: expected a trade object or a list of trades".format(source))
    return list(payload)


def _load_csv(source: Path) -> List[Mapping[str, Any]]:
    with open(source, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("{}: CSV file has no header row".format(source))
        return [
            {key: value for key, value in row.items() if key is not None}
            for row in reader
            if any((value or "").strip() for value in row.values())
        ]


def _load_xlsx(source: Path) -> List[Mapping[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - optional dependency
        raise ImportError(
            "reading .xlsx term sheets requires openpyxl; install it with "
            "'pip install openpyxl' or export the sheet as CSV"
        ) from error
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as error:
        raise ValueError("{}: Excel sheet is empty".format(source)) from error
    payloads = []
    for row in rows:
        if all(value is None for value in row):
            continue
        payloads.append(dict(zip(header, row)))
    return payloads


__all__ = [
    "ObservationRecord",
    "SUPPORTED_PRODUCT_TYPES",
    "TradeTerms",
    "load_terms",
]
